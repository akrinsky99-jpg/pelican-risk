import io
import io
from datetime import datetime
from fastapi.responses import StreamingResponse


def export_excel(record, corrections, customer) -> StreamingResponse:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Audit Trail"

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(fill_type="solid", fgColor="1F4E79")
    label_font  = Font(bold=True)
    center      = Alignment(horizontal="center")
    wrap        = Alignment(wrap_text=True)
    thin_border = Border(bottom=Side(style="thin", color="CCCCCC"))

    def hdr(ws, row, col, value):
        cell = ws.cell(row=row, column=col, value=value)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center

    def label(ws, row, col, value):
        cell = ws.cell(row=row, column=col, value=value)
        cell.font = label_font

    def val(ws, row, col, value):
        cell = ws.cell(row=row, column=col, value=value)
        cell.alignment = wrap
        cell.border = thin_border

    ws.merge_cells("A1:D1")
    title_cell = ws["A1"]
    title_cell.value = "Pelican Risk — Classification Audit Trail"
    title_cell.font = Font(bold=True, size=14, color="1F4E79")
    title_cell.alignment = center

    ws.merge_cells("A2:D2")
    sub = ws["A2"]
    sub.value = f"Institution: {customer.name}  |  Exported: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}"
    sub.alignment = center
    sub.font = Font(italic=True, color="555555")

    fields = [
        ("Business Name",     record.business_name),
        ("Address",           record.address or "—"),
        ("Parish",            record.parish or "—"),
        ("NAICS Code",        record.naics_code or "Unclassified"),
        ("Industry",          record.naics_description or "—"),
        ("Sector",            record.naics_sector or "—"),
        ("Confidence Score",  record.confidence_score),
        ("Confidence Tier",   record.confidence_tier),
        ("Method",            record.method),
        ("Matched Keyword",   record.match_keyword or "none"),
        ("Conflict Detected", "Yes" if record.conflict_detected else "No"),
        ("Needs Review",      "Yes" if record.needs_review else "No"),
        ("Confirmed",         "Yes" if record.confirmed else "No"),
        ("Confirmed By",      record.confirmed_by or "—"),
        ("Confirmed At",      record.confirmed_at.strftime("%Y-%m-%d %H:%M UTC") if record.confirmed_at else "—"),
        ("Classified At",     record.created_at.strftime("%Y-%m-%d %H:%M UTC") if record.created_at else "—"),
        ("Record ID",         record.id),
    ]

    row = 4
    hdr(ws, row, 1, "Field")
    hdr(ws, row, 2, "Value")
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 60

    for field, value in fields:
        row += 1
        label(ws, row, 1, field)
        val(ws, row, 2, value)

    row += 2
    ws.merge_cells(f"A{row}:B{row}")
    hdr(ws, row, 1, "Audit Reasoning")
    ws[f"A{row}"].alignment = Alignment(horizontal="left")

    row += 1
    ws.merge_cells(f"A{row}:B{row}")
    reasoning_cell = ws.cell(row=row, column=1, value=record.reasoning or "No reasoning recorded.")
    reasoning_cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[row].height = 60

    row += 2
    ws.merge_cells(f"A{row}:D{row}")
    hdr(ws, row, 1, f"Correction History ({len(corrections)} corrections)")
    ws[f"A{row}"].alignment = Alignment(horizontal="left")
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 40

    if corrections:
        row += 1
        for col, h in enumerate(["From Code", "To Code", "Corrected By", "Reason", "Date"], start=1):
            cell = ws.cell(row=row, column=col, value=h)
            cell.font = Font(bold=True, color="1F4E79")
        for c in corrections:
            row += 1
            ws.cell(row=row, column=1, value=c.original_naics_code)
            ws.cell(row=row, column=2, value=c.corrected_naics_code)
            ws.cell(row=row, column=3, value=c.corrected_by or "—")
            ws.cell(row=row, column=4, value=c.correction_reason or "—")
            ws.cell(row=row, column=5, value=c.corrected_at.strftime("%Y-%m-%d") if c.corrected_at else "—")
    else:
        row += 1
        ws.merge_cells(f"A{row}:D{row}")
        ws.cell(row=row, column=1, value="No corrections on record.")

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    safe_name = record.business_name.replace(" ", "_").replace("/", "-")[:40]
    filename = f"audit_{safe_name}_{record.id[:8]}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )


def export_pdf(record, corrections, customer) -> StreamingResponse:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.lib import colors
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
    )

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=letter,
        rightMargin=0.75 * inch, leftMargin=0.75 * inch,
        topMargin=0.75 * inch, bottomMargin=0.75 * inch
    )

    styles = getSampleStyleSheet()
    navy  = colors.HexColor("#1F4E79")
    gray  = colors.HexColor("#555555")
    light = colors.HexColor("#EBF2FA")
    warn  = colors.HexColor("#FF6B35")

    title_style   = ParagraphStyle("Title",   parent=styles["Heading1"], textColor=navy, fontSize=16, spaceAfter=4)
    sub_style     = ParagraphStyle("Sub",     parent=styles["Normal"],   textColor=gray, fontSize=9,  spaceAfter=12)
    section_style = ParagraphStyle("Section", parent=styles["Heading2"], textColor=navy, fontSize=11, spaceBefore=14, spaceAfter=6)
    body_style    = ParagraphStyle("Body",    parent=styles["Normal"],   fontSize=9, leading=14)
    warn_style    = ParagraphStyle("Warn",    parent=styles["Normal"],   textColor=warn, fontSize=9)

    story = []

    story.append(Paragraph("Pelican Risk — Classification Audit Trail", title_style))
    story.append(Paragraph(
        f"Institution: {customer.name} &nbsp;&nbsp;|&nbsp;&nbsp; "
        f"Exported: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}",
        sub_style
    ))
    story.append(HRFlowable(width="100%", thickness=1, color=navy, spaceAfter=10))

    story.append(Paragraph("Classification Detail", section_style))

    field_data = [
        ["Business Name",     record.business_name or "—"],
        ["Address",           record.address or "—"],
        ["Parish",            record.parish or "—"],
        ["NAICS Code",        record.naics_code or "Unclassified"],
        ["Industry",          record.naics_description or "—"],
        ["Sector",            record.naics_sector or "—"],
        ["Confidence Score",  str(record.confidence_score)],
        ["Confidence Tier",   (record.confidence_tier or "").upper()],
        ["Method",            record.method or "—"],
        ["Matched Keyword",   record.match_keyword or "none"],
        ["Conflict Detected", "Yes" if record.conflict_detected else "No"],
        ["Needs Review",      "Yes" if record.needs_review else "No"],
        ["Confirmed",         "Yes" if record.confirmed else "No"],
        ["Confirmed By",      record.confirmed_by or "—"],
        ["Confirmed At",      record.confirmed_at.strftime("%Y-%m-%d %H:%M UTC") if record.confirmed_at else "—"],
        ["Classified At",     record.created_at.strftime("%Y-%m-%d %H:%M UTC") if record.created_at else "—"],
        ["Record ID",         record.id],
    ]

    tbl = Table(field_data, colWidths=[1.8 * inch, 5.2 * inch])
    tbl.setStyle(TableStyle([
        ("BACKGROUND",    (0, 0), (0, -1), light),
        ("FONTNAME",      (0, 0), (0, -1), "Helvetica-Bold"),
        ("FONTSIZE",      (0, 0), (-1, -1), 9),
        ("TEXTCOLOR",     (0, 0), (0, -1), navy),
        ("VALIGN",        (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS",(0, 0), (-1, -1), [colors.white, colors.HexColor("#F7FBFF")]),
        ("GRID",          (0, 0), (-1, -1), 0.4, colors.HexColor("#CCCCCC")),
        ("TOPPADDING",    (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("LEFTPADDING",   (0, 0), (-1, -1), 6),
    ]))
    story.append(tbl)

    story.append(Paragraph("Audit Reasoning", section_style))
    story.append(Paragraph(record.reasoning or "No reasoning recorded.", body_style))

    if record.needs_review and not record.confirmed:
        story.append(Spacer(1, 6))
        story.append(Paragraph(
            "⚠ This classification is flagged for human review and has not yet been confirmed.",
            warn_style
        ))

    story.append(Paragraph(f"Correction History ({len(corrections)} corrections)", section_style))

    if corrections:
        corr_data = [["From Code", "To Code", "Corrected By", "Reason", "Date"]]
        for c in corrections:
            corr_data.append([
                c.original_naics_code or "—",
                c.corrected_naics_code,
                c.corrected_by or "—",
                Paragraph(c.correction_reason or "—", body_style),
                c.corrected_at.strftime("%Y-%m-%d") if c.corrected_at else "—",
            ])
        corr_tbl = Table(corr_data, colWidths=[0.8*inch, 0.8*inch, 1.2*inch, 3.2*inch, 1.0*inch])
        corr_tbl.setStyle(TableStyle([
            ("BACKGROUND",    (0, 0), (-1, 0), navy),
            ("TEXTCOLOR",     (0, 0), (-1, 0), colors.white),
            ("FONTNAME",      (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE",      (0, 0), (-1, -1), 8),
            ("ROWBACKGROUNDS",(0, 1), (-1, -1), [colors.white, colors.HexColor("#F7FBFF")]),
            ("GRID",          (0, 0), (-1, -1), 0.4, colors.HexColor("#CCCCCC")),
            ("VALIGN",        (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING",    (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ("LEFTPADDING",   (0, 0), (-1, -1), 6),
        ]))
        story.append(corr_tbl)
    else:
        story.append(Paragraph("No corrections on record.", body_style))

    story.append(Spacer(1, 20))
    story.append(HRFlowable(width="100%", thickness=0.5, color=gray))
    story.append(Spacer(1, 4))
    story.append(Paragraph(
        "This report is generated by Pelican Risk and contains a permanent, unmodified audit record "
        "of all classification decisions and corrections for this business. "
        "It is intended for regulatory examination and internal compliance review.",
        ParagraphStyle("Footer", parent=styles["Normal"], fontSize=7, textColor=gray)
    ))

    doc.build(story)
    buf.seek(0)

    safe_name = record.business_name.replace(" ", "_").replace("/", "-")[:40]
    filename = f"audit_{safe_name}_{record.id[:8]}.pdf"
    return StreamingResponse(
        buf,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )
